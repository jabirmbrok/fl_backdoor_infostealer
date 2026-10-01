"""Put the speaker script into the deck's speaker notes, for PowerPoint's Presenter View.

Reads slides/speaker_script_15min.md ("## Slide N. Title (m:ss, clock m:ss)" sections) and, for
each main slide N, rewrites its notes as:

    Suggested time: m:ss (m:ss of 15:00 at the end of this slide).

    SCRIPT
    <the section's paragraphs; *cues* shown as [cues]>

    Q&A NOTES
    <whatever notes the slide already had, without an older timing line or script>

It also sets each main slide's page number to "N / <main slides>". Rerunning replaces the
earlier timing and script instead of adding another copy. Slide text is not touched otherwise.

Run:  python slides/embed_script.py <script.md> <in.pptx> <out.pptx>
"""
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

CRLF = "\r\n"
TIMING = re.compile(r"^Suggested time: \d+:\d\d \(\d+:\d\d of 15:00 at the end of this slide\)\.\s*")


def sections(md):
    heads = list(re.finditer(r"^## Slide (\d+)\. [^\n(]*\((\d+:\d\d), clock (\d+:\d\d)\)$", md, flags=re.M))
    out = {}
    for i, h in enumerate(heads):
        end = heads[i + 1].start() if i + 1 < len(heads) else len(md)
        body = md[h.end():end].strip()
        body = re.sub(r"^---\s*$", "", body, flags=re.M).strip()
        body = re.sub(r"\*([^*]+)\*", r"[\1]", body)
        paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
        out[int(h.group(1))] = (h.group(2), h.group(3), paras)
    return out


def old_qa(text):
    """Keep only the presenter's own notes: drop a timing line and any earlier script."""
    text = text.replace("\r\n", "\n")
    text = TIMING.sub("", text).strip()
    if text.startswith("SCRIPT"):
        k = text.find("Q&A NOTES")
        text = text[k + len("Q&A NOTES"):].strip() if k >= 0 else ""
    return text


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    paras = tf.paragraphs
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    p = paras[0]
    runs = p.runs
    if not runs:
        run = p.add_run()
    else:
        run = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    # write CR LF straight into <a:t>; the run.text setter would escape CR as _x000D_
    run._r.find(qn("a:t")).text = text


def main(md_path, src, dst):
    md = open(md_path, encoding="utf8").read()
    secs = sections(md)
    prs = Presentation(src)
    slides = list(prs.slides)
    n_main = max(secs)
    assert slides[n_main - 1].shapes and any(
        sh.has_text_frame and sh._element.find(qn("p:txBody")) is not None and sh.text_frame.text.strip() == "Thank you"
        for sh in slides[n_main - 1].shapes), "the script's last section must be the Thank-you slide"
    for n, (slot, clock, paras) in secs.items():
        s = slides[n - 1]
        before = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        qa = old_qa(before)
        text = f"Suggested time: {slot} ({clock} of 15:00 at the end of this slide)." + CRLF + CRLF
        text += "SCRIPT" + CRLF + (CRLF + CRLF).join(paras)
        if qa:
            text += CRLF + CRLF + "Q&A NOTES" + CRLF + qa.replace("\n", CRLF)
        set_notes(s, text)
        if n > 1:
            for sh in s.shapes:
                if sh.has_text_frame and sh._element.find(qn("p:txBody")) is not None and re.fullmatch(r"\d+ / \d+", sh.text_frame.text.strip()):
                    run = sh.text_frame.paragraphs[0].runs[0]
                    run.text = f"{n} / {n_main}"
    prs.save(dst)
    import os
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from add_citations import restore_jpg_default
    restore_jpg_default(dst)
    print(f"script placed in the notes of slides 1-{n_main}; page numbers set to N / {n_main}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
